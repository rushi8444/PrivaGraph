"""Unit and integration tests for Excel (.xlsx) workbook ingestion and querying."""

import io
from datetime import date
from pathlib import Path
import pytest
from openpyxl import Workbook
from httpx import AsyncClient, ASGITransport

from app.ingestion.csv_converter import CSVToOKFConverter
from app.main import app


def create_sample_xlsx_bytes() -> bytes:
    """Create a sample single-sheet Excel workbook in memory."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Payroll"
    ws.append(["Employee Name", "Role", "Department", "Salary", "SSN"])
    ws.append(["Marcus Feldstein", "Staff Engineer", "Engineering", "$196,000", "927-18-3364"])
    ws.append(["Priya Nandakumar", "Senior Analyst", "Finance", "$142,500", "912-04-7731"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def create_multisheet_xlsx_bytes() -> bytes:
    """Create a multi-sheet Excel workbook in memory."""
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Engineering"
    ws1.append(["Employee Name", "Role", "Salary"])
    ws1.append(["David Kim", "DevOps Lead", "$175,000"])

    ws2 = wb.create_sheet(title="Marketing")
    ws2.append(["Employee Name", "Role", "Salary"])
    ws2.append(["Sarah Jenkins", "Content Director", "$160,000"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_converter_single_sheet():
    """Verify that a single-sheet Excel workbook converts into validated OKFDocument."""
    content = create_sample_xlsx_bytes()
    converter = CSVToOKFConverter()
    doc = converter.convert(content, filename="test_payroll.xlsx", department="Finance")

    assert doc.header.document.id.startswith("XLSX-")
    assert doc.header.graph.namespace.startswith("xlsx.")
    assert len(doc.sentences) >= 2
    assert any("Marcus Feldstein" in s for s in doc.sentences)
    assert any("Priya Nandakumar" in s for s in doc.sentences)


def test_xlsx_converter_multi_sheet():
    """Verify that multiple worksheets in an .xlsx workbook are all converted."""
    content = create_multisheet_xlsx_bytes()
    converter = CSVToOKFConverter()
    doc = converter.convert(content, filename="company_roster.xlsx")

    # Sentences should contain records from both sheets
    assert any("David Kim" in s for s in doc.sentences)
    assert any("Sarah Jenkins" in s for s in doc.sentences)


def test_xlsx_converter_empty_workbook_raises_error():
    """Verify that an empty workbook raises ValueError."""
    wb = Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    converter = CSVToOKFConverter()
    with pytest.raises(ValueError, match="no non-empty sheets"):
        converter.convert(buf.getvalue(), filename="empty.xlsx")


@pytest.mark.asyncio
async def test_xlsx_api_ingestion_and_query_flow():
    """Test full API pipeline: upload .xlsx file -> query facts from the graph."""
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        xlsx_bytes = create_sample_xlsx_bytes()
        files = {
            "file": (
                "test_payroll.xlsx",
                xlsx_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        }
        data = {"department": "Finance", "classification": "INTERNAL"}
        ingest_res = await ac.post("/api/v1/documents/ingest", files=files, data=data)

        assert ingest_res.status_code == 200
        res_json = ingest_res.json()
        assert res_json["status"] == "ingested"
        assert res_json["doc_id"].startswith("XLSX-")
        assert res_json["entities_detected"] > 0
        assert res_json["triples_extracted"] > 0

        # Query a fact extracted from the uploaded .xlsx file
        query_res = await ac.post(
            "/api/v1/query",
            json={"query": "What is Marcus Feldstein's salary?"},
        )
        assert query_res.status_code == 200
        answer = query_res.json()["answer"]
        assert "$196,000" in answer


@pytest.mark.asyncio
async def test_xlsx_async_background_ingestion():
    """Test asynchronous background ingestion with .xlsx workbook."""
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        xlsx_bytes = create_multisheet_xlsx_bytes()
        files = {
            "file": (
                "multisheet_team.xlsx",
                xlsx_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        }
        data = {"background": "true"}
        ingest_res = await ac.post("/api/v1/documents/ingest", files=files, data=data)

        assert ingest_res.status_code == 202
        job_id = ingest_res.json()["job_id"]

        # Poll status until complete
        import asyncio
        for _ in range(20):
            status_res = await ac.get(f"/api/v1/documents/jobs/{job_id}")
            if status_res.json()["status"] == "completed":
                break
            await asyncio.sleep(0.1)

        final_job = (await ac.get(f"/api/v1/documents/jobs/{job_id}")).json()
        assert final_job["status"] == "completed"
        assert final_job["result"]["doc_id"].startswith("XLSX-")
